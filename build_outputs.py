"""Create data profile, calculated summaries, notebook, and Excel workbook."""
from pathlib import Path
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
RAW = ROOT / 'data' / 'raw'
OUT = ROOT / 'outputs' / 'analysis_results'
OUT.mkdir(parents=True, exist_ok=True)
for folder in ['excel', 'python', 'outputs/charts']:
    (ROOT / folder).mkdir(parents=True, exist_ok=True)

files = {p.stem: pd.read_csv(p) for p in RAW.glob('*.csv')}
profile = []
for name, df in files.items():
    profile.append({'table': name, 'rows': len(df), 'columns': len(df.columns),
                    'missing_cells': int(df.isna().sum().sum()),
                    'duplicate_rows': int(df.duplicated().sum()),
                    'columns_list': ', '.join(df.columns)})
pd.DataFrame(profile).to_csv(OUT / 'data_profile.csv', index=False)
column_profile=[]
for name, df in files.items():
    for col in df.columns:
        column_profile.append({'table':name,'column':col,'inferred_dtype':str(df[col].dtype),
                               'missing_values':int(df[col].isna().sum()),'date_candidate':('date' in col or 'timestamp' in col)})
pd.DataFrame(column_profile).to_csv(OUT/'column_profile.csv',index=False)
key_columns={
 'olist_orders_dataset':['order_id'], 'olist_order_items_dataset':['order_id','order_item_id'],
 'olist_customers_dataset':['customer_id'], 'olist_products_dataset':['product_id'],
 'olist_sellers_dataset':['seller_id'], 'olist_order_payments_dataset':['order_id','payment_sequential'],
 'olist_order_reviews_dataset':['review_id','order_id'], 'product_category_name_translation':['product_category_name']}
key_profile=[]
for name, cols in key_columns.items():
    df=files[name]
    key_profile.append({'table':name,'key_columns':', '.join(cols),'duplicate_key_rows':int(df.duplicated(cols).sum()),'null_key_rows':int(df[cols].isna().any(axis=1).sum())})
pd.DataFrame(key_profile).to_csv(OUT/'key_profile.csv',index=False)

orders = files['olist_orders_dataset'].copy()
items = files['olist_order_items_dataset'].copy()
customers = files['olist_customers_dataset'].copy()
products = files['olist_products_dataset'].copy()
sellers = files['olist_sellers_dataset'].copy()
payments = files['olist_order_payments_dataset'].copy()
reviews = files['olist_order_reviews_dataset'].copy()
translation = files['product_category_name_translation']
for col in ['order_purchase_timestamp','order_delivered_customer_date','order_estimated_delivery_date']:
    orders[col] = pd.to_datetime(orders[col], errors='coerce')
delivered = orders[orders.order_status.eq('delivered')].copy()
sales = items.merge(delivered[['order_id','customer_id','order_purchase_timestamp','order_delivered_customer_date','order_estimated_delivery_date']], on='order_id', how='inner')
sales['month'] = sales.order_purchase_timestamp.dt.to_period('M').astype(str)
monthly = sales.groupby('month', as_index=False).agg(revenue=('price','sum'), orders=('order_id','nunique'))
category = (sales.merge(products[['product_id','product_category_name']], on='product_id', how='left')
            .merge(translation, on='product_category_name', how='left'))
category['category'] = category.product_category_name_english.fillna(category.product_category_name).fillna('unknown')
category_summary = category.groupby('category',as_index=False).agg(revenue=('price','sum'),items=('order_id','size'),orders=('order_id','nunique')).sort_values('revenue',ascending=False)
state = (sales.merge(customers[['customer_id','customer_unique_id','customer_state']],on='customer_id',how='left')
         .groupby('customer_state',as_index=False).agg(revenue=('price','sum'),orders=('order_id','nunique'))
         .sort_values('revenue',ascending=False))
customer_order_counts = delivered.merge(customers[['customer_id','customer_unique_id']],on='customer_id',how='left').groupby('customer_unique_id').order_id.nunique()
customer_sales = sales.merge(customers[['customer_id','customer_unique_id']],on='customer_id',how='left').groupby('customer_unique_id',as_index=False).price.sum().rename(columns={'price':'revenue'}).sort_values('revenue',ascending=False)
review_by_order = reviews.groupby('order_id',as_index=False).review_score.mean()
order_reviews = delivered[['order_id','order_delivered_customer_date','order_estimated_delivery_date']].merge(review_by_order,on='order_id',how='inner')
order_reviews['late'] = order_reviews.order_delivered_customer_date > order_reviews.order_estimated_delivery_date
order_reviews['delivery_group'] = order_reviews.late.map({True:'Late',False:'On/before estimate'})
review_dist = reviews.groupby('review_score',as_index=False).size().rename(columns={'size':'reviews'})
delivery = delivered.dropna(subset=['order_delivered_customer_date','order_estimated_delivery_date']).copy()
delivery['delivery_days']=(delivery.order_delivered_customer_date-delivery.order_purchase_timestamp).dt.total_seconds()/86400
delivery['late']=delivery.order_delivered_customer_date>delivery.order_estimated_delivery_date
payment_mix = payments.groupby('payment_type',as_index=False).agg(payment_value=('payment_value','sum'),orders=('order_id','nunique')).sort_values('payment_value',ascending=False)
seller_summary=(sales.groupby('seller_id',as_index=False).agg(revenue=('price','sum'),orders=('order_id','nunique'))
 .merge(sellers[['seller_id','seller_state']],on='seller_id',how='left').sort_values('revenue',ascending=False))
review_summary=order_reviews.groupby('delivery_group',as_index=False).agg(orders=('order_id','nunique'),avg_score=('review_score','mean'))

metrics = {
 'Delivered item revenue (freight excluded)': sales.price.sum(),
 'Delivered orders': delivered.order_id.nunique(),
 'Unique customers with delivered orders': int(customer_order_counts.index.nunique()),
 'Repeat customers': int((customer_order_counts>1).sum()),
 'Repeat customer share (%)': float((customer_order_counts>1).mean()*100),
 'Average order value': float(sales.groupby('order_id').price.sum().mean()),
 'Average review score (order average)': float(review_by_order.review_score.mean()),
 'Average delivery days': float(delivery.delivery_days.mean()),
 'Delivered orders after estimate (%)': float(delivery.late.mean()*100),
 'Late review average score': float(order_reviews.loc[order_reviews.late,'review_score'].mean()),
 'On-time review average score': float(order_reviews.loc[~order_reviews.late,'review_score'].mean()),
}
pd.DataFrame([{'metric':k,'value':v} for k,v in metrics.items()]).to_csv(OUT/'kpi_summary.csv',index=False)
monthly.to_csv(OUT/'monthly_sales.csv',index=False)
category_summary.to_csv(OUT/'category_sales.csv',index=False)
state.to_csv(OUT/'state_sales.csv',index=False)
review_dist.to_csv(OUT/'review_distribution.csv',index=False)
payment_mix.to_csv(OUT/'payment_mix.csv',index=False)
review_summary.to_csv(OUT/'delivery_review_comparison.csv',index=False)
insights = [
 {'finding':f"Delivered item revenue totals {metrics['Delivered item revenue (freight excluded)']:,.2f} across {metrics['Delivered orders']:,} delivered orders.", 'business_impact':'This provides the sales baseline using item prices, without mixing freight or payment value into revenue.', 'recommendation':'Track item revenue and freight separately in the sales dashboard.'},
 {'finding':f"Average delivered order value is {metrics['Average order value']:,.2f}.", 'business_impact':'AOV gives a simple view of the typical delivered basket value.', 'recommendation':'Monitor AOV by month and category before testing basket-building promotions.'},
 {'finding':f"{metrics['Repeat customer share (%)']:.2f}% of delivered customers ordered more than once ({metrics['Repeat customers']:,} repeat customers).", 'business_impact':'Most customer identities in this measure are single-delivery buyers; retention is a potential area to examine.', 'recommendation':'Compare repeat share by acquisition month and customer state, then assess targeted retention offers.'},
 {'finding':f"The highest-revenue category is {category_summary.iloc[0].category} ({category_summary.iloc[0].revenue:,.2f} in item sales).", 'business_impact':'Category sales are concentrated unevenly, so category mix affects overall sales performance.', 'recommendation':'Review inventory availability and seller coverage for the leading categories.'},
 {'finding':f"The leading customer state by delivered item revenue is {state.iloc[0].customer_state} ({state.iloc[0].revenue:,.2f}).", 'business_impact':'Sales demand is not evenly distributed across locations.', 'recommendation':'Use state-level demand when planning campaigns and fulfillment capacity.'},
 {'finding':f"Average purchase-to-customer delivery took {metrics['Average delivery days']:.2f} days; {metrics['Delivered orders after estimate (%)']:.2f}% of completed orders arrived after the estimate.", 'business_impact':'Late delivery is a direct service measure and may influence review outcomes.', 'recommendation':'Prioritize states and sellers with both high late share and meaningful order volume.'},
 {'finding':f"The average order-level review score was {metrics['Average review score (order average)']:.2f} out of 5.", 'business_impact':'The score summarizes customer satisfaction, while individual categories and order experiences can differ.', 'recommendation':'Review low-scoring categories and pair them with delivery results before selecting operational actions.'},
 {'finding':f"The most-used payment type by payment value is {payment_mix.iloc[0].payment_type}.", 'business_impact':'Payment mix helps prioritize the checkout methods customers use most.', 'recommendation':'Monitor payment method mix and transaction failures separately; payment value is not item revenue.'},
]
pd.DataFrame(insights).to_csv(OUT/'business_insights.csv',index=False)

# Workbook with source-backed summaries, formulas for KPIs, and editable charts.
wb=Workbook(); ws=wb.active; ws.title='KPI Summary'
ws.append(['E-Commerce Sales & Operations Analytics','Value'])
for row in [('Delivered item revenue',"=SUM('Sales Analysis'!B2:B1000)"),('Delivered orders',"=SUM('Sales Analysis'!C2:C1000)"),('Average order value',"=IFERROR(B2/B3,0)"),('Unique delivered customers',int(customer_order_counts.size)),('Repeat delivered customers',int((customer_order_counts>1).sum()))]: ws.append(list(row))
def sheet(name, headers, rows):
    sh=wb.create_sheet(name); sh.append(headers)
    for r in rows: sh.append([None if pd.isna(x) else (x.item() if hasattr(x,'item') else x) for x in r])
    sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
    for cell in sh[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
    for col in sh.columns:
        width=min(38,max(12,max(len(str(c.value or '')) for c in list(col)[:200])+2))
        sh.column_dimensions[get_column_letter(col[0].column)].width=width
    return sh
# Monthly summaries keep the workbook quick to open and easy to scan.
sheet('Sales Analysis',['month','revenue','orders'],monthly.itertuples(index=False,name=None))
sheet('Category Analysis',['category','revenue','items','orders'],category_summary.itertuples(index=False,name=None))
customer_segments=pd.DataFrame([{'segment':'One-time','customers':int((customer_order_counts==1).sum())},{'segment':'Repeat','customers':int((customer_order_counts>1).sum())}])
sheet('Customer Analysis',['segment','customers'],customer_segments.itertuples(index=False,name=None))
sheet('Seller Analysis',['seller_id','seller_state','revenue','orders'],seller_summary[['seller_id','seller_state','revenue','orders']].itertuples(index=False,name=None))
delivery_state=delivery.merge(customers[['customer_id','customer_state']],on='customer_id').groupby('customer_state',as_index=False).agg(avg_delivery_days=('delivery_days','mean'),late_pct=('late','mean'),orders=('order_id','nunique'))
delivery_state['late_pct']*=100
sheet('Delivery Analysis',['customer_state','avg_delivery_days','late_pct','orders'],delivery_state.itertuples(index=False,name=None))
sheet('Monthly Sales',['month','revenue','orders'],monthly.itertuples(index=False,name=None))
sheet('Review Analysis',['score','reviews'],review_dist.itertuples(index=False,name=None))
sheet('Payment Analysis',['payment_type','payment_value','orders'],payment_mix.itertuples(index=False,name=None))
# Store review distribution and detailed summaries; spreadsheet charts are beginner-friendly.
for sh in wb.worksheets:
    for cell in sh[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
catsh=wb['Category Analysis']; chart=BarChart(); chart.title='Revenue by category'; chart.y_axis.title='Revenue'; chart.x_axis.title='Category'; chart.add_data(Reference(catsh,min_col=2,min_row=1,max_row=min(11,catsh.max_row)),titles_from_data=True); chart.set_categories(Reference(catsh,min_col=1,min_row=2,max_row=min(11,catsh.max_row))); wb['KPI Summary'].add_chart(chart,'D2')
msh=wb['Monthly Sales']; lc=LineChart(); lc.title='Monthly delivered item revenue'; lc.add_data(Reference(msh,min_col=2,min_row=1,max_row=msh.max_row),titles_from_data=True); lc.set_categories(Reference(msh,min_col=1,min_row=2,max_row=msh.max_row)); wb['KPI Summary'].add_chart(lc,'D20')
wb.save(ROOT/'excel'/'ecommerce_analysis.xlsx')

notebook_cells = [
('md','# E-Commerce Sales & Operations Analytics\n\nA practical walkthrough of Olist orders, item sales, customers, categories, delivery, reviews and payments. Item revenue is the sum of item prices and is not profit.'),
('code',"from pathlib import Path\nimport pandas as pd\nimport numpy as np\nimport matplotlib.pyplot as plt\nimport seaborn as sns\n\nDATA_DIR = Path('../data/raw')\norders = pd.read_csv(DATA_DIR / 'olist_orders_dataset.csv')\nitems = pd.read_csv(DATA_DIR / 'olist_order_items_dataset.csv')\ncustomers = pd.read_csv(DATA_DIR / 'olist_customers_dataset.csv')\nproducts = pd.read_csv(DATA_DIR / 'olist_products_dataset.csv')\nsellers = pd.read_csv(DATA_DIR / 'olist_sellers_dataset.csv')\nreviews = pd.read_csv(DATA_DIR / 'olist_order_reviews_dataset.csv')\npayments = pd.read_csv(DATA_DIR / 'olist_order_payments_dataset.csv')\ntranslation = pd.read_csv(DATA_DIR / 'product_category_name_translation.csv')\nprint('orders:', orders.shape, 'items:', items.shape, 'customers:', customers.shape)"),
('md','## 1. Inspect structure and data quality\n\nThe source tables have different grains. Orders is one row per order; order items has multiple rows per order. Do not join item, payment and review detail together before aggregation.'),
('code',"tables = {'orders': orders, 'items': items, 'customers': customers, 'products': products, 'sellers': sellers, 'reviews': reviews, 'payments': payments}\nprofile = pd.DataFrame([{'table': name, 'rows': len(df), 'columns': len(df.columns), 'duplicate_rows': df.duplicated().sum(), 'missing_cells': df.isna().sum().sum()} for name, df in tables.items()])\ndisplay(profile)\ndisplay(orders.isna().sum().sort_values(ascending=False).head(10))\nprint('Order key duplicates:', orders.order_id.duplicated().sum())\nprint('Order item key duplicates:', items.duplicated(['order_id', 'order_item_id']).sum())"),
('md','## 2. Prepare dates and delivered sales\n\nDelivered orders are used for realized sales and delivery analysis. Non-delivered orders remain relevant to order-status analysis. Freight is kept separate from item revenue.'),
('code',"date_cols = ['order_purchase_timestamp', 'order_approved_at', 'order_delivered_carrier_date', 'order_delivered_customer_date', 'order_estimated_delivery_date']\nfor col in date_cols:\n    orders[col] = pd.to_datetime(orders[col], errors='coerce')\ndelivered = orders.loc[orders.order_status.eq('delivered')].copy()\nsales = items.merge(delivered[['order_id', 'customer_id', 'order_purchase_timestamp', 'order_delivered_customer_date', 'order_estimated_delivery_date']], on='order_id', how='inner', validate='many_to_one')\nsales['month'] = sales.order_purchase_timestamp.dt.to_period('M')\nprint('Delivered item revenue:', round(sales.price.sum(), 2))\nprint('Delivered orders:', sales.order_id.nunique())\nprint('Average order value:', round(sales.groupby('order_id').price.sum().mean(), 2))"),
('md','## 3. Monthly sales\n\nMonthly revenue is grouped by purchase month. The chart only includes delivered orders, so recent months may be incomplete if the extract ends before fulfillment catches up.'),
('code',"monthly = sales.groupby('month').agg(revenue=('price', 'sum'), orders=('order_id', 'nunique')).reset_index()\nfig, ax = plt.subplots(figsize=(11, 4))\nsns.lineplot(data=monthly, x=monthly.month.astype(str), y='revenue', marker='o', ax=ax)\nax.set(title='Monthly delivered item revenue', xlabel='Purchase month', ylabel='Revenue'); ax.tick_params(axis='x', rotation=60); plt.tight_layout()\nmonthly.plot(x=monthly.month.astype(str), y='orders', kind='bar', figsize=(11, 4), legend=False, title='Delivered orders by purchase month'); plt.ylabel('Orders'); plt.tight_layout()"),
('md','## 4. Category and geographic analysis\n\nTranslate category labels to English where available. Customer state is used for geographic sales because it is directly present in the customer table.'),
('code',"category = sales.merge(products[['product_id', 'product_category_name']], on='product_id', how='left').merge(translation, on='product_category_name', how='left')\ncategory['category'] = category.product_category_name_english.fillna(category.product_category_name).fillna('unknown')\ncategory_sales = category.groupby('category').price.sum().sort_values(ascending=False)\ncategory_sales.head(12).sort_values().plot.barh(figsize=(8, 5), title='Top categories by delivered item revenue'); plt.xlabel('Revenue'); plt.tight_layout()\nstate_sales = sales.merge(customers[['customer_id', 'customer_state']], on='customer_id').groupby('customer_state').price.sum().sort_values(ascending=False)\nstate_sales.head(10).sort_values().plot.barh(figsize=(7, 4), title='Top customer states by revenue'); plt.xlabel('Revenue'); plt.tight_layout()"),
('md','## 5. Customer behavior\n\nUse `customer_unique_id` as a person-level identifier. `customer_id` is linked to an individual order record in this dataset.'),
('code',"customer_orders = delivered.merge(customers[['customer_id', 'customer_unique_id']], on='customer_id').groupby('customer_unique_id').order_id.nunique()\nrepeat_customers = int((customer_orders > 1).sum())\nprint('Unique delivered customers:', customer_orders.size)\nprint('Repeat customers:', repeat_customers)\nprint('Repeat share (%):', round((customer_orders > 1).mean() * 100, 2))"),
('md','## 6. Reviews and delivery\n\nReview text is optional; the numeric score supports satisfaction analysis. Compare at order level because an order can have more than one review record.'),
('code',"reviews.review_creation_date = pd.to_datetime(reviews.review_creation_date, errors='coerce')\nreviews.review_score.value_counts().sort_index().plot.bar(figsize=(6, 4), title='Review score distribution'); plt.xlabel('Score'); plt.ylabel('Review records'); plt.tight_layout()\ndelivery = delivered.dropna(subset=['order_delivered_customer_date']).copy()\ndelivery['delivery_days'] = (delivery.order_delivered_customer_date - delivery.order_purchase_timestamp).dt.total_seconds() / 86400\ndelivery['late'] = delivery.order_delivered_customer_date > delivery.order_estimated_delivery_date\nsns.histplot(data=delivery, x='delivery_days', bins=40); plt.title('Purchase-to-customer delivery time'); plt.xlabel('Days'); plt.tight_layout()\nprint('Average delivery days:', round(delivery.delivery_days.mean(), 2))\nprint('Late share (%):', round(delivery.late.mean() * 100, 2))\norder_score = reviews.groupby('order_id', as_index=False).review_score.mean()\nreview_delivery = delivery[['order_id', 'late']].merge(order_score, on='order_id')\nprint(review_delivery.groupby('late').review_score.mean())"),
('md','## 7. Payments and validation\n\nPayment value is transaction value, not product revenue. Keep it separate from item sales because installment rows and payment adjustments can make the measures differ.'),
('code',"display(payments.groupby('payment_type').agg(payment_value=('payment_value', 'sum'), orders=('order_id', 'nunique')).sort_values('payment_value', ascending=False))\nprint('Item revenue:', round(sales.price.sum(), 2))\nprint('Freight on delivered items:', round(sales.freight_value.sum(), 2))\nprint('Payment value on delivered orders:', round(payments[payments.order_id.isin(delivered.order_id)].payment_value.sum(), 2))"),
]
def source(kind, text):
    return {'cell_type':'markdown' if kind=='md' else 'code','metadata':{},'source':text.splitlines(True), **({'execution_count':None,'outputs':[]} if kind=='code' else {})}
nb={'cells':[source(k,t) for k,t in notebook_cells], 'metadata':{'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'},'language_info':{'name':'python','version':'3'}},'nbformat':4,'nbformat_minor':5}
(ROOT/'python'/'ecommerce_eda.ipynb').write_text(json.dumps(nb,indent=2),encoding='utf-8')
print(json.dumps(metrics,indent=2))
